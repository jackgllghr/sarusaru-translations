#!/usr/bin/env python3
# /// script
# requires-python = ">=3.9"
# dependencies = ["openpyxl"]
# ///
"""Translate the Monkey Book in a spreadsheet.

  scripts/monkey_sheet.py export [-o monkeys.xlsx] [--iso game.iso]
  scripts/monkey_sheet.py import monkeys.xlsx

`export` writes one row per monkey (name, kind, description and the square,
circle and cross actions) from english-gameplay-v1.json, plus a sheet for the
book's stage header and catch counter. `import` writes the edited English back
into the catalog.

The catalog keeps each distinct kind and action once, under the first record
that uses it. Which entry every record shows is kept in monkey-book-records.json
(IDs only), so exporting needs no game files. In the sheet, a record that
shares another's kind or action shows it greyed out as a formula pointing at
that record's cell: edit the text in the record that owns it.

`--iso` reads your copy of the game to refresh monkey-book-records.json and
adds the original Japanese beside each column. A sheet with Japanese text is
for your own use: do not commit or share it (*.xlsx is git-ignored).

Changed text is marked origin human_translation; set a row's Status to
reviewed once a person has checked it. Needs openpyxl (`uv run` installs it).
"""
import argparse
import hashlib
import json
import os
from pathlib import Path
import struct
import sys

HERE = Path(__file__).resolve().parents[1]
CATALOG = HERE / 'english-gameplay-v1.json'
RECORDS = HERE / 'monkey-book-records.json'
FIELDS = ('name', 'kind', 'description', 'square', 'circle', 'cross')
SHARED = ('kind', 'square', 'circle', 'cross')
HEADERS = {'name': 'Name', 'kind': 'Kind', 'description': 'Description',
           'square': 'Square action', 'circle': 'Circle action', 'cross': 'Cross action'}
WIDTHS = {'name': 16, 'kind': 22, 'description': 60, 'square': 26, 'circle': 26, 'cross': 26}
STATUSES = ('draft', 'reviewed')
MONKEYS, BOOK_UI = 'Monkeys', 'Book UI'


def load_catalog(path):
    text = path.read_text(encoding='utf-8')
    catalog = json.loads(text)
    entries = {e['id']: e for e in catalog['strings'] + catalog.get('monkey_book', [])}
    return catalog, entries


def save_catalog(path, catalog):
    path.write_text(json.dumps(catalog, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')


def record_key(ident):
    return f'{ident:03d}'


def entry_record(entry_id):
    """'monkey.record_012.kind' -> '012'."""
    return entry_id.split('.')[1].removeprefix('record_')


# Reading the game, only with --iso.

def recomp_tools():
    recomp = Path(os.environ.get('SARUSARU_RECOMP', HERE.parent / 'bigmission-recomp'))
    tools = recomp / 'profiles/sarusaru/tools'
    if not (tools / 'prepare_game.py').is_file():
        sys.exit(f'monkey_sheet: no bigmission-recomp checkout at {recomp} (set SARUSARU_RECOMP)')
    return recomp / 'profiles/sarusaru', tools


def read_game(iso_path):
    """Per record: (entry ID, Japanese source) for each field, in table order."""
    profile, tools = recomp_tools()
    sys.path.insert(0, str(tools))
    from prepare_game import IsoImage
    supported = json.loads((profile / 'config/supported_game.json').read_text(encoding='utf-8'))
    metadata = json.loads((profile / 'resources/translations/extraction-context.json').read_text(encoding='utf-8'))
    with iso_path.open('rb') as stream:
        boot = IsoImage(stream).file(supported['executable_path'], 4_000_000)
    if hashlib.sha256(boot).hexdigest() != supported['executable_sha256']:
        sys.exit('monkey_sheet: unsupported executable; only UCJS10060 version 1.00 is supported')

    def cstring(offset):
        end = boot.index(b'\0', offset)
        return boot[offset:end]

    table, book = metadata['monkey_names'], metadata['monkey_book']
    base, owners, records = int(table['offset'], 0), {}, {}
    for index in range(table['count']):
        at = base + index * table['stride']
        key = record_key(struct.unpack_from('<I', boot, at)[0])
        name = struct.unpack_from('<I', boot, at + table['pointer_offset'])[0] + 0x74
        raws = {'name': cstring(name)}
        for field, offset in book['fields'].items():
            raws[field] = cstring(struct.unpack_from('<I', boot, at + offset)[0] + book['pointer_bias'])
        record = {}
        for field in FIELDS:
            raw = raws[field]
            if not raw:
                record[field] = None
                continue
            # Same rule as the recomp's monkey_book.sources: one entry per
            # distinct text, named after the first record that uses it.
            entry_id = owners.setdefault((field, raw) if field in SHARED else (field, key),
                                         f'monkey.record_{key}.{field}')
            record[field] = (entry_id, raw.decode('cp932'), hashlib.sha256(raw + b'\0').hexdigest())
        records[key] = record
    ui = {f'monkey_book.{k}': cstring(int(v, 0)).decode('cp932') for k, v in book['ui'].items()}
    ui['monkey_book.caught_once'] = ui['monkey_book.caught']
    return records, ui


def write_records(records):
    data = {
        'description': 'Which catalog entry each Monkey Book record shows; written by scripts/monkey_sheet.py --iso.',
        'records': {key: {field: value and value[0] for field, value in record.items()}
                    for key, record in sorted(records.items())},
    }
    RECORDS.write_text(json.dumps(data, indent=2) + '\n', encoding='utf-8')


def load_records():
    if not RECORDS.is_file():
        sys.exit(f'monkey_sheet: {RECORDS.name} is missing; run export once with --iso')
    return json.loads(RECORDS.read_text(encoding='utf-8'))['records']


def check_records(records, entries, sources=None):
    problems = []
    for key, record in records.items():
        for field, entry_id in record.items():
            if entry_id is None:
                continue
            if entry_id not in entries:
                problems.append(f'record {key} {field}: {entry_id} is not in the catalog')
            elif sources and entries[entry_id]['source_sha256'] != sources[key][field][2]:
                problems.append(f'{entry_id}: the source hash does not match the game')
    used = {entry_id for record in records.values() for entry_id in record.values()}
    for entry_id in entries:
        if entry_id.startswith('monkey.record_') and entry_id not in used:
            problems.append(f'{entry_id} is not shown by any record')
    if problems:
        sys.exit('monkey_sheet: the catalog and the record map disagree:\n  ' + '\n  '.join(problems))


def row_status(entries, ids):
    return 'reviewed' if ids and all(entries[i]['review_status'] == 'reviewed' for i in ids) else 'draft'


def owned_ids(key, record):
    return [entry_id for entry_id in record.values() if entry_id and entry_record(entry_id) == key]


# Export.

def export(args):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    catalog, entries = load_catalog(args.catalog)
    sources, ui_sources = None, {}
    if args.iso:
        sources, ui_sources = read_game(args.iso)
        records = {key: {field: value and value[0] for field, value in record.items()}
                   for key, record in sources.items()}
        check_records(records, entries, sources)
        write_records(sources)
        print(f'monkey_sheet: wrote {RECORDS.name}')
    else:
        records = load_records()
        check_records(records, entries)

    bold, header_fill = Font(bold=True), PatternFill('solid', fgColor='DDE4EE')
    shared_font, japanese_font = Font(color='8A8A8A', italic=True), Font(color='5A6B80')
    empty_fill = PatternFill('solid', fgColor='EEEEEE')
    wrap = Alignment(wrap_text=True, vertical='top')
    workbook = Workbook()

    sheet = workbook.active
    sheet.title = MONKEYS
    columns = ['No.']
    for field in FIELDS:
        if sources:
            columns.append(f'{HEADERS[field]} (Japanese)')
        columns.append(HEADERS[field])
    columns.append('Status')
    sheet.append(columns)
    column = {name: index + 1 for index, name in enumerate(columns)}
    for cell in sheet[1]:
        cell.font, cell.fill = bold, header_fill
    keys = sorted(records)
    row_of = {key: row for row, key in enumerate(keys, start=2)}
    sharers = {}
    for key in keys:
        for field in SHARED:
            entry_id = records[key][field]
            if entry_id and entry_record(entry_id) != key:
                sharers.setdefault(entry_id, []).append(key)

    for key in keys:
        row, record = row_of[key], records[key]
        sheet.cell(row, column['No.'], key)
        for field in FIELDS:
            cell = sheet.cell(row, column[HEADERS[field]])
            cell.alignment = wrap
            if sources and sources[key][field]:
                source = sheet.cell(row, column[f'{HEADERS[field]} (Japanese)'], sources[key][field][1])
                source.font, source.alignment = japanese_font, wrap
            entry_id = record[field]
            if entry_id is None:
                cell.fill = empty_fill
            elif entry_record(entry_id) == key:
                cell.value = entries[entry_id]['text']
                if entry_id in sharers:
                    others = ', '.join(sharers[entry_id])
                    cell.comment = Comment(f'Also shown by {len(sharers[entry_id])} other records: {others}', 'monkey_sheet')
            else:
                owner = row_of[entry_record(entry_id)]
                cell.value = f'={get_column_letter(cell.column)}{owner}'
                cell.font = shared_font
        sheet.cell(row, column['Status'], row_status(entries, owned_ids(key, record)))

    status = DataValidation(type='list', formula1='"' + ','.join(STATUSES) + '"', allow_blank=False)
    sheet.add_data_validation(status)
    letter = get_column_letter(column['Status'])
    status.add(f'{letter}2:{letter}{len(keys) + 1}')
    sheet.column_dimensions['A'].width = 6
    for field in FIELDS:
        sheet.column_dimensions[get_column_letter(column[HEADERS[field]])].width = WIDTHS[field]
        if sources:
            sheet.column_dimensions[get_column_letter(column[f'{HEADERS[field]} (Japanese)'])].width = WIDTHS[field]
    sheet.column_dimensions[letter].width = 11
    sheet.freeze_panes = 'C2' if sources else 'B2'
    sheet.auto_filter.ref = sheet.dimensions

    ui = workbook.create_sheet(BOOK_UI)
    ui.append(['ID', 'Japanese', 'English', 'Status'] if sources else ['ID', 'English', 'Status'])
    for cell in ui[1]:
        cell.font, cell.fill = bold, header_fill
    for entry in catalog.get('monkey_book', []):
        if entry['id'].startswith('monkey_book.'):
            row = [entry['id']] + ([ui_sources.get(entry['id'], '')] if sources else [])
            ui.append(row + [entry['text'], entry['review_status']])
    ui.add_data_validation(ui_status := DataValidation(type='list', formula1='"' + ','.join(STATUSES) + '"'))
    ui_status.add(f'{get_column_letter(ui.max_column)}2:{get_column_letter(ui.max_column)}{ui.max_row}')
    ui.column_dimensions['A'].width = 26
    for letter in 'BCD':
        ui.column_dimensions[letter].width = 22
    ui.append([])
    ui.append(['{n} is replaced by the number.'])

    workbook.save(args.output)
    print(f'monkey_sheet: wrote {args.output} ({len(keys)} monkeys)')


# Import.

def cell_text(value):
    if value is None:
        return None
    return str(value)


def check_text(entry_id, text, problems):
    if not text or not text.strip():
        problems.append(f'{entry_id}: the English text is empty')
    elif any(not 32 <= ord(c) <= 126 for c in text):
        problems.append(f'{entry_id}: the game draws printable ASCII only: {text!r}')


def apply(entry, text, status, changes):
    if text is not None and text != entry['text']:
        entry['text'] = text
        entry['origin'] = 'human_translation'
        changes.append(entry['id'])
    if status is not None and status != entry['review_status']:
        entry['review_status'] = status
        if entry['id'] not in changes:
            changes.append(entry['id'])


def import_(args):
    from openpyxl import load_workbook

    catalog, entries = load_catalog(args.catalog)
    records = load_records()
    check_records(records, entries)
    workbook = load_workbook(args.sheet)
    problems, changes, texts, statuses = [], [], {}, {}

    sheet = workbook[MONKEYS]
    header = {cell.value: cell.column for cell in sheet[1] if cell.value}
    missing = [h for h in ['No.', *HEADERS.values(), 'Status'] if h not in header]
    if missing:
        sys.exit(f'monkey_sheet: the {MONKEYS} sheet has no column {", ".join(missing)}')
    seen = set()
    for row in sheet.iter_rows(min_row=2):
        key = row[header['No.'] - 1].value
        if key is None:
            continue
        key = record_key(int(key))
        if key not in records or key in seen:
            problems.append(f'row {row[0].row}: unknown or repeated record {key}')
            continue
        seen.add(key)
        record = records[key]
        status = cell_text(row[header['Status'] - 1].value)
        if status not in STATUSES:
            problems.append(f'record {key}: Status must be draft or reviewed, not {status!r}')
            status = None
        elif status != row_status(entries, owned_ids(key, record)):
            for entry_id in owned_ids(key, record):
                statuses[entry_id] = status
        for field in FIELDS:
            entry_id, value = record[field], cell_text(row[header[HEADERS[field]] - 1].value)
            if entry_id is None:
                if value:
                    problems.append(f'record {key} has no {field} in the game; clear its cell')
            elif entry_record(entry_id) == key:
                check_text(entry_id, value, problems)
                texts[entry_id] = value
            elif value is not None and not value.startswith('='):
                owner = entry_record(entry_id)
                problems.append(f'record {key} {field} is shared with record {owner}; '
                                f'edit it in row {owner} and restore the formula')
    if seen != set(records):
        problems.append(f'the {MONKEYS} sheet is missing records {", ".join(sorted(set(records) - seen))}')

    if BOOK_UI in workbook.sheetnames:
        ui = workbook[BOOK_UI]
        ui_header = {cell.value: cell.column for cell in ui[1] if cell.value}
        for row in ui.iter_rows(min_row=2):
            entry_id = cell_text(row[ui_header['ID'] - 1].value)
            if not entry_id or not entry_id.startswith('monkey_book.'):
                continue
            if entry_id not in entries:
                problems.append(f'{BOOK_UI}: unknown ID {entry_id}')
                continue
            text = cell_text(row[ui_header['English'] - 1].value)
            check_text(entry_id, text, problems)
            if '{n}' in entries[entry_id]['text'] and text and '{n}' not in text:
                problems.append(f'{entry_id}: keep the {{n}} placeholder')
            texts[entry_id] = text
            status = cell_text(row[ui_header['Status'] - 1].value)
            if status not in STATUSES:
                problems.append(f'{entry_id}: Status must be draft or reviewed, not {status!r}')
            else:
                statuses[entry_id] = status

    if problems:
        sys.exit('monkey_sheet: nothing imported:\n  ' + '\n  '.join(problems))
    for entry_id in sorted(set(texts) | set(statuses)):
        apply(entries[entry_id], texts.get(entry_id), statuses.get(entry_id), changes)
    if changes:
        save_catalog(args.catalog, catalog)
    print(f'monkey_sheet: updated {len(changes)} entries in {args.catalog.name}')
    for entry_id in changes:
        print(f'  {entry_id}: {entries[entry_id]["text"]!r} ({entries[entry_id]["review_status"]})')


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--catalog', type=Path, default=CATALOG)
    commands = parser.add_subparsers(dest='command', required=True)
    out = commands.add_parser('export', help='write the spreadsheet')
    out.add_argument('-o', '--output', type=Path, default=HERE / 'monkeys.xlsx')
    out.add_argument('--iso', type=Path, help='your copy of the game: refreshes the record map and adds Japanese')
    into = commands.add_parser('import', help='write the spreadsheet back into the catalog')
    into.add_argument('sheet', type=Path)
    args = parser.parse_args()
    export(args) if args.command == 'export' else import_(args)


if __name__ == '__main__':
    main()
